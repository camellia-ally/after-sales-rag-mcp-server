"""
将 markdown 文档转换为 PDF 和 Word 格式，用于测试 RAG 文档加载器
"""
import os
from pathlib import Path
from fpdf import FPDF
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


class UTF8PDF(FPDF):
    """支持中文的 PDF 类"""
    def __init__(self):
        super().__init__()
        # 添加中文字体支持
        try:
            self.add_font('SimHei', '', 'C:/Windows/Fonts/simhei.ttf')
            self.add_font('SimHei', 'B', 'C:/Windows/Fonts/simhei.ttf')
            self.font_family = 'SimHei'
        except Exception as e:
            print(f"警告: 无法加载中文字体 ({e})，使用默认字体")
            self.font_family = 'Arial'

    def chapter_title(self, title):
        self.set_font(self.font_family, 'B', 16)
        self.cell(0, 10, title, 0, 1, 'L')
        self.ln(4)

    def chapter_body(self, body):
        self.set_font(self.font_family, '', 11)
        self.multi_cell(0, 6, body)
        self.ln()


def markdown_to_pdf(md_path, pdf_path):
    """将 markdown 转换为 PDF"""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    pdf = UTF8PDF()
    pdf.add_page()

    lines = content.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            pdf.ln(2)
            continue

        if line.startswith('# '):
            # 一级标题
            pdf.set_font(pdf.font_family, 'B', 16)
            pdf.cell(0, 10, line[2:], 0, 1, 'L')
            pdf.ln(2)
        elif line.startswith('## '):
            # 二级标题
            pdf.set_font(pdf.font_family, 'B', 14)
            pdf.cell(0, 8, line[3:], 0, 1, 'L')
            pdf.ln(2)
        elif line.startswith('### '):
            # 三级标题
            pdf.set_font(pdf.font_family, 'B', 12)
            pdf.cell(0, 7, line[4:], 0, 1, 'L')
            pdf.ln(1)
        elif line.startswith('#### '):
            # 四级标题
            pdf.set_font(pdf.font_family, 'B', 11)
            pdf.cell(0, 6, line[5:], 0, 1, 'L')
            pdf.ln(1)
        elif line.startswith('- ') or line.startswith('* '):
            # 列表项
            pdf.set_font(pdf.font_family, '', 10)
            pdf.cell(10, 5, '•', 0, 0, 'L')
            pdf.multi_cell(0, 5, line[2:])
        elif line.startswith('**') and line.endswith('**'):
            # 粗体
            pdf.set_font(pdf.font_family, 'B', 10)
            pdf.multi_cell(0, 5, line.strip('*'))
        else:
            # 普通文本
            pdf.set_font(pdf.font_family, '', 10)
            pdf.multi_cell(0, 5, line)

    pdf.output(pdf_path)
    print(f"[OK] 生成 PDF: {pdf_path}")


def markdown_to_word(md_path, docx_path):
    """将 markdown 转换为 Word"""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    lines = content.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            continue

        if line.startswith('# '):
            # 一级标题
            heading = doc.add_heading(line[2:], level=1)
            heading.runs[0].font.name = '微软雅黑'
            heading.runs[0].font.size = Pt(16)
        elif line.startswith('## '):
            # 二级标题
            heading = doc.add_heading(line[3:], level=2)
            heading.runs[0].font.name = '微软雅黑'
            heading.runs[0].font.size = Pt(14)
        elif line.startswith('### '):
            # 三级标题
            heading = doc.add_heading(line[4:], level=3)
            heading.runs[0].font.name = '微软雅黑'
            heading.runs[0].font.size = Pt(12)
        elif line.startswith('#### '):
            # 四级标题
            heading = doc.add_heading(line[5:], level=4)
            heading.runs[0].font.name = '微软雅黑'
            heading.runs[0].font.size = Pt(11)
        elif line.startswith('- ') or line.startswith('* '):
            # 列表项
            p = doc.add_paragraph(line[2:], style='List Bullet')
            p.runs[0].font.name = '宋体'
            p.runs[0].font.size = Pt(10.5)
        elif line.startswith('**') and line.endswith('**'):
            # 粗体
            p = doc.add_paragraph()
            run = p.add_run(line.strip('*'))
            run.bold = True
            run.font.name = '宋体'
            run.font.size = Pt(10.5)
        else:
            # 普通段落
            p = doc.add_paragraph(line)
            p.runs[0].font.name = '宋体'
            p.runs[0].font.size = Pt(10.5)

    doc.save(docx_path)
    print(f"[OK] 生成 Word: {docx_path}")


def main():
    """批量转换所有 markdown 文档"""
    base_dir = Path(__file__).parent / 'data' / 'customer_service'

    # 遍历所有 markdown 文件
    md_files = list(base_dir.rglob('*.md'))

    if not md_files:
        print("未找到 markdown 文件")
        return

    print(f"找到 {len(md_files)} 个 markdown 文件\n")

    for md_file in md_files:
        try:
            # 生成相对路径以保持目录结构
            relative_path = md_file.relative_to(base_dir)

            # PDF 输出路径
            pdf_path = base_dir / relative_path.parent / f"{md_file.stem}.pdf"
            pdf_path.parent.mkdir(parents=True, exist_ok=True)

            # Word 输出路径
            docx_path = base_dir / relative_path.parent / f"{md_file.stem}.docx"
            docx_path.parent.mkdir(parents=True, exist_ok=True)

            print(f"处理: {relative_path}")
            markdown_to_pdf(str(md_file), str(pdf_path))
            markdown_to_word(str(md_file), str(docx_path))
            print()

        except Exception as e:
            print(f"[ERROR] 转换失败 {md_file.name}: {e}\n")

    print("转换完成！")


if __name__ == '__main__':
    main()
